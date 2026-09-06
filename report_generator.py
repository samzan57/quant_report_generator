# report_generator.py
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from groq import Groq
import json
import os
from datetime import datetime
from config import GROQ_API_KEY, GROQ_MODEL, OUTPUT_DIR

client = Groq(api_key=GROQ_API_KEY)

# Couleurs
DARK_BG = RGBColor(13, 17, 23)
CARD_BG = RGBColor(22, 27, 34)
ACCENT = RGBColor(0, 212, 255)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(139, 148, 158)
GREEN = RGBColor(0, 255, 136)
RED = RGBColor(255, 68, 68)
GOLD = RGBColor(255, 215, 0)

def add_background(slide, prs):
    """Ajoute un fond sombre à la slide"""
    background = slide.shapes.add_shape(
        1, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    return background

def add_text_box(slide, text, left, top, width, height,
                 font_size=18, bold=False, color=WHITE,
                 align=PP_ALIGN.LEFT, wrap=True):
    """Ajoute une zone de texte stylée"""
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return txBox

def add_card(slide, left, top, width, height, color=CARD_BG):
    """Ajoute une carte avec fond"""
    card = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = color
    card.line.color.rgb = ACCENT
    card.line.width = Pt(0.5)
    return card

def create_title_slide(prs, analysis, asset_name):
    """Slide de titre"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    # Ligne décorative
    line = slide.shapes.add_shape(
        1, Inches(0), Inches(3.5), prs.slide_width, Pt(3)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    # Titre principal
    add_text_box(slide, "RAPPORT D'ANALYSE QUANTITATIVE",
                 0.5, 1.2, 12, 1, font_size=32, bold=True,
                 color=ACCENT, align=PP_ALIGN.CENTER)

    # Sous-titre
    add_text_box(slide, asset_name.upper(),
                 0.5, 2.3, 12, 0.8, font_size=22, bold=True,
                 color=WHITE, align=PP_ALIGN.CENTER)

    # Type de données
    add_text_box(slide, analysis.get('type_donnees', ''),
                 0.5, 3.8, 12, 0.6, font_size=14,
                 color=GRAY, align=PP_ALIGN.CENTER)

    # Période
    add_text_box(slide, f"Période : {analysis.get('periode', 'N/A')}",
                 0.5, 4.4, 12, 0.6, font_size=13,
                 color=GRAY, align=PP_ALIGN.CENTER)

    # Date du rapport
    date_str = datetime.now().strftime("%d %B %Y")
    add_text_box(slide, f"Généré le {date_str}",
                 0.5, 5.8, 12, 0.5, font_size=11,
                 color=GRAY, align=PP_ALIGN.CENTER)

    return slide

def create_summary_slide(prs, analysis):
    """Slide de résumé exécutif"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, "RÉSUMÉ EXÉCUTIF", 0.3, 0.2, 12, 0.7,
                 font_size=20, bold=True, color=ACCENT)

    # Ligne
    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.95), Inches(12.1), Pt(1.5)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    # Résumé IA
    resume = analysis.get('resume', 'Analyse en cours...')
    add_card(slide, 0.3, 1.1, 12.1, 1.5)
    add_text_box(slide, resume, 0.5, 1.2, 11.7, 1.3,
                 font_size=13, color=WHITE)

    # Actifs détectés
    actifs = ", ".join(analysis.get('actifs_detectes', ['N/A']))
    add_text_box(slide, "ACTIFS ANALYSÉS", 0.3, 2.8, 5.8, 0.4,
                 font_size=11, bold=True, color=ACCENT)
    add_card(slide, 0.3, 3.2, 5.8, 0.8)
    add_text_box(slide, actifs, 0.5, 3.25, 5.6, 0.7,
                 font_size=12, color=WHITE)

    # Qualité des données
    qualite = analysis.get('qualite_donnees', 'N/A')
    qualite_color = GREEN if qualite == 'bonne' else (
        GOLD if qualite == 'moyenne' else RED)
    add_text_box(slide, "QUALITÉ DES DONNÉES", 6.5, 2.8, 5.8, 0.4,
                 font_size=11, bold=True, color=ACCENT)
    add_card(slide, 6.5, 3.2, 5.8, 0.8)
    add_text_box(slide, qualite.upper(), 6.7, 3.25, 5.6, 0.7,
                 font_size=14, bold=True, color=qualite_color)

    # Analyses recommandées
    analyses = analysis.get('analyses_recommandees', [])
    add_text_box(slide, "ANALYSES RÉALISÉES", 0.3, 4.2, 12.1, 0.4,
                 font_size=11, bold=True, color=ACCENT)
    add_card(slide, 0.3, 4.6, 12.1, 2.0)
    analyses_text = "  •  ".join(analyses[:6])
    add_text_box(slide, analyses_text, 0.5, 4.7, 11.7, 1.8,
                 font_size=11, color=WHITE)

    return slide

def create_metrics_slide(prs, metrics, asset_name):
    """Slide des métriques quantitatives"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, f"MÉTRIQUES QUANTITATIVES — {asset_name.upper()}",
                 0.3, 0.2, 12, 0.7, font_size=18, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.95), Inches(12.1), Pt(1.5)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    # Métriques à afficher
    metrics_display = [
        ("Rendement Total", f"{metrics.get('rendement_total', 0):.2f}%",
         GREEN if metrics.get('rendement_total', 0) > 0 else RED),
        ("Rendement Ann.", f"{metrics.get('rendement_annualise', 0):.2f}%",
         GREEN if metrics.get('rendement_annualise', 0) > 0 else RED),
        ("Volatilité Ann.", f"{metrics.get('volatilite_annualisee', 0):.2f}%", GOLD),
        ("Sharpe Ratio", f"{metrics.get('sharpe_ratio', 0):.3f}",
         GREEN if metrics.get('sharpe_ratio', 0) > 1 else GOLD),
        ("Sortino Ratio", f"{metrics.get('sortino_ratio', 0):.3f}",
         GREEN if metrics.get('sortino_ratio', 0) > 1 else GOLD),
        ("Max Drawdown", f"{metrics.get('max_drawdown', 0):.2f}%", RED),
        ("VaR 95%", f"{metrics.get('var_95', 0):.2f}%", RED),
        ("CVaR 95%", f"{metrics.get('cvar_95', 0):.2f}%", RED),
        ("Win Rate", f"{metrics.get('win_rate', 0):.1f}%",
         GREEN if metrics.get('win_rate', 0) > 50 else RED),
        ("Beta", f"{metrics.get('beta', 0):.3f}", WHITE),
        ("Skewness", f"{metrics.get('skewness', 0):.3f}", WHITE),
        ("Kurtosis", f"{metrics.get('kurtosis', 0):.3f}", WHITE),
    ]

    cols = 4
    rows = 3
    card_w = 2.9
    card_h = 1.4
    start_x = 0.3
    start_y = 1.1

    for i, (label, value, color) in enumerate(metrics_display):
        col = i % cols
        row = i // cols
        x = start_x + col * (card_w + 0.15)
        y = start_y + row * (card_h + 0.15)

        add_card(slide, x, y, card_w, card_h)
        add_text_box(slide, value, x, y + 0.2, card_w, 0.7,
                     font_size=22, bold=True, color=color,
                     align=PP_ALIGN.CENTER)
        add_text_box(slide, label, x, y + 0.85, card_w, 0.45,
                     font_size=10, color=GRAY, align=PP_ALIGN.CENTER)

    return slide

def create_chart_slide(prs, chart_path, title):
    """Slide avec un graphique"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, title.upper(), 0.3, 0.15, 12, 0.6,
                 font_size=16, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.8), Inches(12.1), Pt(1.5)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    if os.path.exists(chart_path):
        slide.shapes.add_picture(
            chart_path,
            Inches(0.3), Inches(1.0),
            Inches(12.1), Inches(5.7)
        )

    return slide

def generate_ai_commentary(metrics, analysis):
    """Génère un commentaire IA sur les métriques"""
    prompt = f"""
Tu es un analyste quantitatif expert. Voici les métriques d'un actif financier :

{json.dumps(metrics, indent=2)}

Contexte : {analysis.get('resume', '')}

Génère un commentaire professionnel en français de 4-5 phrases qui :
1. Évalue la performance globale
2. Commente le ratio risque/rendement
3. Identifie les points forts et faibles
4. Donne une conclusion sur la qualité de cet actif

Sois précis, professionnel et factuel. Utilise les chiffres.
"""
    response = client.chat.completions.create(
        model=GROQ_MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.3
    )
    return response.choices[0].message.content.strip()

def create_commentary_slide(prs, metrics, analysis, asset_name):
    """Slide avec commentaire IA"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, "ANALYSE & COMMENTAIRES IA",
                 0.3, 0.2, 12, 0.7, font_size=20, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.95), Inches(12.1), Pt(1.5)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    print("Génération du commentaire IA...")
    commentary = generate_ai_commentary(metrics, analysis)

    add_card(slide, 0.3, 1.1, 12.1, 3.5)
    add_text_box(slide, commentary, 0.5, 1.2, 11.7, 3.3,
                 font_size=13, color=WHITE)

    # Points clés
    sharpe = metrics.get('sharpe_ratio', 0)
    drawdown = metrics.get('max_drawdown', 0)
    win_rate = metrics.get('win_rate', 0)

    points = [
        ("SHARPE RATIO",
         "Excellent" if sharpe > 2 else "Bon" if sharpe > 1 else "Faible",
         GREEN if sharpe > 1 else RED),
        ("DRAWDOWN MAX",
         "Faible" if abs(drawdown) < 10 else "Modéré" if abs(drawdown) < 25 else "Élevé",
         GREEN if abs(drawdown) < 10 else RED),
        ("WIN RATE",
         "Bon" if win_rate > 55 else "Moyen" if win_rate > 45 else "Faible",
         GREEN if win_rate > 55 else RED),
    ]

    for i, (label, verdict, color) in enumerate(points):
        x = 0.3 + i * 4.1
        add_card(slide, x, 4.8, 3.9, 1.5)
        add_text_box(slide, label, x, 4.9, 3.9, 0.5,
                     font_size=10, bold=True, color=GRAY,
                     align=PP_ALIGN.CENTER)
        add_text_box(slide, verdict, x, 5.3, 3.9, 0.7,
                     font_size=18, bold=True, color=color,
                     align=PP_ALIGN.CENTER)

    return slide

def create_conclusion_slide(prs, analysis):
    """Slide de conclusion"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    # Ligne décorative
    line = slide.shapes.add_shape(
        1, Inches(0), Inches(3.2), prs.slide_width, Pt(3)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    add_text_box(slide, "CONCLUSION", 0.5, 1.0, 12, 0.8,
                 font_size=28, bold=True, color=ACCENT,
                 align=PP_ALIGN.CENTER)

    add_text_box(slide,
                 "Rapport généré automatiquement par le système d'analyse quantitative",
                 0.5, 2.0, 12, 0.6, font_size=14,
                 color=GRAY, align=PP_ALIGN.CENTER)

    add_text_box(slide, "Powered by Groq AI & Python Quant Engine",
                 0.5, 3.5, 12, 0.6, font_size=13,
                 color=WHITE, align=PP_ALIGN.CENTER)

    date_str = datetime.now().strftime("%d %B %Y — %H:%M")
    add_text_box(slide, date_str, 0.5, 4.5, 12, 0.5,
                 font_size=12, color=GRAY, align=PP_ALIGN.CENTER)

    return slide

def create_stress_summary_slide(prs, stress_results, asset_name):
    """Slide résumé stress testing"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, f"STRESS TESTING — {asset_name.upper()}",
                 0.3, 0.2, 12, 0.7, font_size=18, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.95), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    crises = stress_results.get('crises_historiques', {})
    resume = stress_results.get('resume', {})

    # Résumé
    if resume:
        add_card(slide, 0.3, 1.1, 12.1, 0.9)
        pire = resume.get('pire_scenario', 0)
        crise = resume.get('crise_la_plus_impactante', 'N/A')
        add_text_box(
            slide,
            f"Pire scénario : {pire:.1f}%  |  "
            f"Crise la plus impactante : {crise}",
            0.5, 1.2, 11.7, 0.7, font_size=12, color=RED
        )

    # Tableau des crises
    y_start = 2.1
    cols = [0.3, 4.5, 7.0, 9.5, 11.5]
    headers = ['Crise', 'Durée', 'Choc Hist.', 'Perte Sim.', 'VaR Stress']

    for i, (header, x) in enumerate(zip(headers, cols)):
        add_text_box(slide, header, x, y_start, 2.5, 0.4,
                     font_size=10, bold=True, color=ACCENT)

    for i, (crisis_name, crisis_data) in enumerate(
            list(crises.items())[:5]):
        y = y_start + 0.5 + i * 0.65
        add_card(slide, 0.3, y, 12.1, 0.55)
        row_data = [
            crisis_name.split('(')[0].strip(),
            f"{crisis_data['duree_jours']}j",
            f"{crisis_data['choc_marche_historique']:.1f}%",
            f"{crisis_data['perte_simulee']:.1f}%",
            f"{crisis_data['var_sous_stress']:.2f}%",
        ]
        for j, (val, x) in enumerate(zip(row_data, cols)):
            color = RED if j >= 2 and '-' in str(val) else WHITE
            add_text_box(slide, val, x, y + 0.08, 2.3, 0.4,
                         font_size=10, color=color)

    return slide

def create_factor_summary_slide(prs, factor_results, asset_name):
    """Slide résumé factor model"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, f"MODÈLES FACTORIELS — {asset_name.upper()}",
                 0.3, 0.2, 12, 0.7, font_size=18, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.95), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    y = 1.1
    for key, data in factor_results.items():
        add_card(slide, 0.3, y, 12.1, 1.6)
        add_text_box(slide, data.get('modele', key).upper(),
                     0.5, y + 0.1, 4.0, 0.4,
                     font_size=11, bold=True, color=ACCENT)

        alpha = data.get('alpha_annualise', 0)
        beta = data.get('beta_marche', data.get('beta', 0))
        r2 = data.get('r_squared', 0)

        alpha_color = GREEN if alpha > 0 else RED
        add_text_box(slide, f"Alpha : {alpha:.2f}%",
                     0.5, y + 0.55, 3.5, 0.4,
                     font_size=13, bold=True, color=alpha_color)
        add_text_box(slide, f"Beta Marché : {beta:.3f}",
                     4.2, y + 0.55, 3.5, 0.4,
                     font_size=13, color=WHITE)
        add_text_box(slide, f"R² : {r2:.3f}",
                     8.0, y + 0.55, 3.5, 0.4,
                     font_size=13, color=WHITE)

        interp = data.get('interpretation', {})
        interp_text = "  |  ".join(
            [f"{k}: {v}" for k, v in list(interp.items())[:3]]
        )
        add_text_box(slide, interp_text,
                     0.5, y + 1.0, 11.7, 0.45,
                     font_size=9, color=GRAY)
        y += 1.8

    return slide

def create_monte_carlo_slide(prs, mc_data, asset_name):
    """Slide résumé Monte Carlo"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, f"MONTE CARLO — {asset_name.upper()}",
                 0.3, 0.2, 12, 0.7, font_size=18, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.95), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    stats = mc_data.get('stats', {})

    metrics = [
        ("Prix Actuel", f"{stats.get('prix_actuel', 0):.2f}", WHITE),
        ("Prix Médian (1an)", f"{stats.get('prix_median', 0):.2f}", GOLD),
        ("Rendement Médian", f"{stats.get('rendement_median', 0):.1f}%",
         GREEN if stats.get('rendement_median', 0) > 0 else RED),
        ("Probabilité Hausse", f"{stats.get('prob_hausse', 0):.1f}%",
         GREEN if stats.get('prob_hausse', 0) > 50 else RED),
        ("Prob. +10%", f"{stats.get('prob_hausse_10', 0):.1f}%", GREEN),
        ("Prob. -10%", f"{stats.get('prob_baisse_10', 0):.1f}%", RED),
        ("VaR 95% MC", f"{stats.get('var_95', 0):.2f}", RED),
        ("Volatilité Ann.", f"{stats.get('sigma_annualise', 0):.1f}%", GOLD),
    ]

    cols = 4
    card_w = 2.9
    card_h = 1.4
    for i, (label, value, color) in enumerate(metrics):
        col = i % cols
        row = i // cols
        x = 0.3 + col * (card_w + 0.15)
        y = 1.1 + row * (card_h + 0.15)
        add_card(slide, x, y, card_w, card_h)
        add_text_box(slide, value, x, y + 0.2, card_w, 0.7,
                     font_size=22, bold=True, color=color,
                     align=PP_ALIGN.CENTER)
        add_text_box(slide, label, x, y + 0.85, card_w, 0.45,
                     font_size=10, color=GRAY, align=PP_ALIGN.CENTER)

    # Paramètres simulation
    add_text_box(
        slide,
        f"Basé sur {stats.get('n_simulations', 0)} simulations "
        f"sur {stats.get('n_days', 252)} jours — "
        f"Modèle GBM (Geometric Brownian Motion)",
        0.3, 6.8, 12.1, 0.5, font_size=10, color=GRAY,
        align=PP_ALIGN.CENTER
    )

    return slide

def create_backtest_slide(prs, backtest_results, asset_name):
    """Slide résumé backtesting"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, f"BACKTESTING — {asset_name.upper()}",
                 0.3, 0.2, 12, 0.7, font_size=18, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.95), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    # Headers
    cols_x = [0.3, 3.2, 5.5, 7.5, 9.5, 11.3]
    headers = ['Stratégie', 'Rendement', 'Volatilité',
               'Sharpe', 'Drawdown', 'Win Rate']

    y_header = 1.1
    for header, x in zip(headers, cols_x):
        add_text_box(slide, header, x, y_header, 2.5, 0.4,
                     font_size=10, bold=True, color=ACCENT)

    # Données
    for i, (strat_name, strat_data) in enumerate(
            backtest_results.items()):
        y = y_header + 0.5 + i * 0.75
        metrics = strat_data.get('metrics', {})

        is_best_sharpe = metrics.get('sharpe_ratio', 0) == max(
            d['metrics'].get('sharpe_ratio', 0)
            for d in backtest_results.values()
        )
        card_color = RGBColor(0, 40, 20) if is_best_sharpe else CARD_BG
        add_card(slide, 0.3, y, 12.5, 0.65)

        ret = metrics.get('rendement_total', 0)
        sharpe = metrics.get('sharpe_ratio', 0)
        dd = metrics.get('max_drawdown', 0)
        wr = metrics.get('win_rate', 0)
        vol = metrics.get('volatilite', 0)

        row_vals = [
            (strat_name, WHITE),
            (f"{ret:.1f}%",
             GREEN if ret > 0 else RED),
            (f"{vol:.1f}%", GOLD),
            (f"{sharpe:.3f}",
             GREEN if sharpe > 1 else GOLD if sharpe > 0 else RED),
            (f"{dd:.1f}%", RED),
            (f"{wr:.1f}%",
             GREEN if wr > 55 else GOLD if wr > 45 else RED),
        ]

        for (val, color), x in zip(row_vals, cols_x):
            add_text_box(slide, val, x, y + 0.12, 2.3, 0.45,
                         font_size=11, color=color,
                         bold=(strat_name == val and is_best_sharpe))

    # Note
    add_text_box(
        slide,
        "★ Meilleur Sharpe Ratio  |  "
        "Les performances passées ne garantissent pas les résultats futurs",
        0.3, 6.8, 12.1, 0.45, font_size=9, color=GRAY,
        align=PP_ALIGN.CENTER
    )

    return slide

def create_signals_slide(prs, signals_data, asset_name):
    """Slide des signaux de trading"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    final = signals_data.get('signal_final', {})
    signal = final.get('signal_final', 'NEUTRE')
    conviction = final.get('conviction_finale', 50)

    signal_color = (GREEN if signal == 'ACHETER'
                    else RED if signal == 'VENDRE'
                    else GOLD)

    add_text_box(slide,
                 f"SIGNAUX DE TRADING — {asset_name.upper()}",
                 0.3, 0.15, 12, 0.6,
                 font_size=18, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.8), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    # Signal principal
    add_card(slide, 0.3, 1.0, 12.1, 1.3)
    add_text_box(slide, f"SIGNAL FINAL : {signal}",
                 0.3, 1.05, 8.0, 0.7,
                 font_size=26, bold=True, color=signal_color,
                 align=PP_ALIGN.CENTER)
    add_text_box(slide, f"Conviction : {conviction:.0f}%",
                 8.3, 1.05, 4.0, 0.7,
                 font_size=20, bold=True, color=signal_color,
                 align=PP_ALIGN.CENTER)
    add_text_box(slide,
                 final.get('interpretation', ''),
                 0.5, 1.7, 12.0, 0.5,
                 font_size=11, color=WHITE,
                 align=PP_ALIGN.CENTER)

    # Signaux individuels
    signaux = signals_data.get('signaux_individuels', [])
    cols_x = [0.3, 2.55, 4.8, 7.05, 9.3, 11.55]
    card_w = 2.1

    for i, (sig, x) in enumerate(zip(signaux, cols_x)):
        if sig is None:
            continue
        s_color = (GREEN if sig['signal'] == 'ACHETER'
                   else RED if sig['signal'] == 'VENDRE'
                   else GOLD)
        add_card(slide, x, 2.5, card_w, 2.8)
        add_text_box(slide, sig['indicateur'],
                     x, 2.6, card_w, 0.45,
                     font_size=9, bold=True, color=ACCENT,
                     align=PP_ALIGN.CENTER)
        add_text_box(slide, sig['signal'],
                     x, 3.05, card_w, 0.6,
                     font_size=16, bold=True, color=s_color,
                     align=PP_ALIGN.CENTER)
        add_text_box(slide, f"{sig['conviction']:.0f}% conviction",
                     x, 3.65, card_w, 0.4,
                     font_size=10, color=WHITE,
                     align=PP_ALIGN.CENTER)
        add_text_box(slide, sig['raison'],
                     x + 0.05, 4.1, card_w - 0.1, 1.0,
                     font_size=8, color=GRAY,
                     align=PP_ALIGN.CENTER)

    # Scores
    add_card(slide, 0.3, 5.5, 5.8, 1.2)
    add_text_box(slide,
                 f"🟢 Score Achat : {final.get('score_achat', 0):.1f}%",
                 0.5, 5.6, 5.4, 0.5,
                 font_size=14, bold=True, color=GREEN)
    add_text_box(slide,
                 f"Signaux haussiers : {final.get('nb_signaux_achat', 0)}",
                 0.5, 6.0, 5.4, 0.5,
                 font_size=11, color=GRAY)

    add_card(slide, 6.6, 5.5, 5.8, 1.2)
    add_text_box(slide,
                 f"🔴 Score Vente : {final.get('score_vente', 0):.1f}%",
                 6.8, 5.6, 5.4, 0.5,
                 font_size=14, bold=True, color=RED)
    add_text_box(slide,
                 f"Signaux baissiers : {final.get('nb_signaux_vente', 0)}",
                 6.8, 6.0, 5.4, 0.5,
                 font_size=11, color=GRAY)

    add_text_box(slide,
                 "⚠ Ces signaux sont indicatifs et basés sur l'analyse technique. "
                 "Ils ne constituent pas un conseil en investissement.",
                 0.3, 6.9, 12.1, 0.45,
                 font_size=9, color=GRAY,
                 align=PP_ALIGN.CENTER)

    return slide

def create_bs_greeks_slide(prs, bs_results, asset_name):
    """Slide résumé Black-Scholes et Greeks"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide, f"BLACK-SCHOLES & GREEKS — {asset_name.upper()}",
                 0.3, 0.15, 12, 0.6,
                 font_size=18, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.8), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    params = bs_results.get('parametres', {})
    add_card(slide, 0.3, 1.0, 12.1, 0.8)
    add_text_box(slide,
                 f"S = {params.get('S', 0):.2f}  |  "
                 f"σ = {params.get('sigma', 0):.1f}%  |  "
                 f"r = {params.get('r', 0):.1f}%  |  "
                 f"Modèle : Black-Scholes (1973)",
                 0.5, 1.1, 11.7, 0.6,
                 font_size=13, color=WHITE,
                 align=PP_ALIGN.CENTER)

    # Tableau des options par maturité
    headers_x = [0.3, 2.3, 4.1, 5.9, 7.7, 9.5, 11.3]
    headers = ['Maturité', 'Call Prix', 'Put Prix',
               'Delta C', 'Gamma', 'Vega', 'Theta C']

    y_header = 2.0
    for header, x in zip(headers, headers_x):
        add_text_box(slide, header, x, y_header, 2.0, 0.4,
                     font_size=10, bold=True, color=ACCENT)

    for i, (mat_name, mat_data) in enumerate(
            bs_results.get('options', {}).items()):
        y = y_header + 0.5 + i * 0.75
        cg = mat_data['call']['greeks']
        pg = mat_data['put']['greeks']
        add_card(slide, 0.3, y, 12.5, 0.65)

        row_vals = [
            (mat_name, WHITE),
            (f"{cg['prix']:.4f}", GREEN),
            (f"{pg['prix']:.4f}", RED),
            (f"{cg['delta']:.4f}", WHITE),
            (f"{cg['gamma']:.6f}", GOLD),
            (f"{cg['vega']:.4f}", WHITE),
            (f"{cg['theta']:.4f}", RED),
        ]

        for (val, color), x in zip(row_vals, headers_x):
            add_text_box(slide, val, x, y + 0.12, 1.9, 0.45,
                         font_size=10, color=color)

    # Probabilité d'exercice ATM 3 mois
    analyse = bs_results.get('analyse_principale', {})
    if analyse:
        call_prob = analyse['call'].get('prob_exercise', 0)
        put_prob = analyse['put'].get('prob_exercise', 0)
        break_even_c = analyse['call'].get('break_even', 0)
        break_even_p = analyse['put'].get('break_even', 0)

        add_card(slide, 0.3, 5.5, 5.8, 1.2)
        add_text_box(slide, "ATM Call (3 mois)",
                     0.5, 5.6, 5.4, 0.4,
                     font_size=10, bold=True, color=GREEN)
        add_text_box(slide,
                     f"Prob. exercice : {call_prob:.1f}%  |  "
                     f"Break-even : {break_even_c:.2f}",
                     0.5, 6.0, 5.4, 0.5,
                     font_size=11, color=WHITE)

        add_card(slide, 6.6, 5.5, 5.8, 1.2)
        add_text_box(slide, "ATM Put (3 mois)",
                     6.8, 5.6, 5.4, 0.4,
                     font_size=10, bold=True, color=RED)
        add_text_box(slide,
                     f"Prob. exercice : {put_prob:.1f}%  |  "
                     f"Break-even : {break_even_p:.2f}",
                     6.8, 6.0, 5.4, 0.5,
                     font_size=11, color=WHITE)

    return slide

def create_mc_options_slide(prs, mc_options, asset_name):
    """Slide résumé Monte Carlo Options"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide,
                 f"PRICING MONTE CARLO — OPTIONS EXOTIQUES — {asset_name.upper()}",
                 0.3, 0.15, 12, 0.6,
                 font_size=16, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.8), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    params = mc_options.get('parametres', {})
    add_card(slide, 0.3, 1.0, 12.1, 0.7)
    add_text_box(slide,
                 f"S = {params.get('S', 0):.2f}  |  "
                 f"K = {params.get('K', 0):.2f}  |  "
                 f"σ = {params.get('sigma', 0):.1f}%  |  "
                 f"T = 3 mois  |  "
                 f"r = {params.get('r', 0):.1f}%",
                 0.5, 1.1, 11.7, 0.5,
                 font_size=12, color=WHITE,
                 align=PP_ALIGN.CENTER)

    # Tableau comparatif
    headers_x = [0.3, 4.5, 6.8, 9.1, 11.0]
    headers = ["Type d'Option", 'Prix MC',
               'vs Vanille', 'Réduction %', 'Interprétation']

    y_h = 1.9
    for header, x in zip(headers, headers_x):
        add_text_box(slide, header, x, y_h, 2.5, 0.4,
                     font_size=10, bold=True, color=ACCENT)

    vanilla_price = mc_options['vanille']['call']['prix_bs']
    comparaison = mc_options.get('comparaison', {})

    options_data = [
        ('Call Vanille (BS)',
         mc_options['vanille']['call']['prix_bs'],
         0, 0, 'Référence'),
        ('Call Vanille (MC)',
         mc_options['vanille']['call']['prix_mc'],
         mc_options['vanille']['call']['prix_mc'] - vanilla_price,
         mc_options['vanille']['call']['erreur_pct'],
         f"Erreur vs BS: {mc_options['vanille']['call']['erreur_pct']:.2f}%"),
        ('Call Asiatique',
         mc_options['asiatiques']['call_arith']['prix_mc'],
         mc_options['asiatiques']['call_arith']['prix_mc'] - vanilla_price,
         mc_options['asiatiques']['call_arith']['reduction_pct'],
         'Moyenne lisse la vol'),
        ('Call Barrière KO',
         mc_options['barrieres']['knock_out_up']['prix_mc'],
         mc_options['barrieres']['knock_out_up']['prix_mc'] - vanilla_price,
         -mc_options['barrieres']['knock_out_up']['reduction_pct'],
         f"KO: {mc_options['barrieres']['knock_out_up']['pct_trajectoires_ko']:.1f}% paths"),
        ('Call Lookback',
         mc_options['lookback']['call_fixed']['prix_mc'],
         mc_options['lookback']['call_fixed']['prix_mc'] - vanilla_price,
         mc_options['lookback']['call_fixed']['premium_pct'],
         'Prix max garanti'),
        ('Digital Call',
         mc_options['digitales']['digital_call']['prix_mc'],
         mc_options['digitales']['digital_call']['prix_mc'] - vanilla_price,
         0,
         f"Prob ITM: {mc_options['digitales']['digital_call']['prob_in_the_money']:.1f}%"),
    ]

    for i, (name, price, diff, pct, interp) in enumerate(options_data):
        y = y_h + 0.5 + i * 0.65
        add_card(slide, 0.3, y, 12.5, 0.6)

        diff_color = (GREEN if diff >= 0 else RED)
        row = [
            (name, WHITE),
            (f"{price:.4f}", WHITE),
            (f"{diff:+.4f}", diff_color),
            (f"{pct:+.2f}%" if pct != 0 else '—', diff_color),
            (interp, GRAY),
        ]
        for (val, color), x in zip(row, headers_x):
            add_text_box(slide, val, x, y + 0.1, 2.4, 0.45,
                         font_size=10, color=color)

    return slide

def create_heston_slide(prs, heston_results, asset_name):
    """Slide résumé Heston"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide,
                 f"MODÈLE DE HESTON — VOLATILITÉ STOCHASTIQUE — {asset_name.upper()}",
                 0.3, 0.15, 12, 0.6,
                 font_size=16, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.8), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    params = heston_results['parametres_heston']
    marche = heston_results['parametres_marche']
    pricing = heston_results.get('pricing_atm', {})

    # Paramètres Heston
    param_cards = [
        ('κ Vitesse retour', f"{params['kappa']:.4f}", ACCENT),
        ('θ Var. long terme', f"{params['theta']:.4f}", GREEN),
        ('ξ Vol de Vol', f"{params['xi']:.4f}", GOLD),
        ('ρ Corrélation', f"{params['rho']:.4f}", RGBColor(155, 89, 182)),
        ('v₀ Var. initiale', f"{params['v0']:.4f}", RGBColor(255, 107, 53)),
        ("Feller",
         "✓" if params['condition_feller'] else "✗",
         GREEN if params['condition_feller'] else RED),
    ]

    card_w = 1.95
    for i, (label, value, color) in enumerate(param_cards):
        x = 0.3 + i * (card_w + 0.1)
        add_card(slide, x, 1.0, card_w, 1.3)
        add_text_box(slide, value, x, 1.1, card_w, 0.6,
                     font_size=18, bold=True, color=color,
                     align=PP_ALIGN.CENTER)
        add_text_box(slide, label, x, 1.65, card_w, 0.45,
                     font_size=9, color=GRAY,
                     align=PP_ALIGN.CENTER)

    # Comparaison pricing
    add_text_box(slide, "COMPARAISON PRICING ATM (3 MOIS)",
                 0.3, 2.5, 12.1, 0.4,
                 font_size=11, bold=True, color=ACCENT)

    add_card(slide, 0.3, 2.95, 5.8, 1.5)
    add_text_box(slide, "BLACK-SCHOLES",
                 0.5, 3.05, 5.4, 0.4,
                 font_size=11, bold=True, color=ACCENT,
                 align=PP_ALIGN.CENTER)
    add_text_box(slide, f"{pricing.get('bs_call', 0):.4f}",
                 0.5, 3.45, 5.4, 0.6,
                 font_size=24, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER)
    add_text_box(slide, f"Vol constante : {marche['sigma_bs']:.1f}%",
                 0.5, 3.95, 5.4, 0.4,
                 font_size=10, color=GRAY,
                 align=PP_ALIGN.CENTER)

    add_card(slide, 6.6, 2.95, 5.8, 1.5)
    add_text_box(slide, "HESTON",
                 6.8, 3.05, 5.4, 0.4,
                 font_size=11, bold=True, color=GREEN,
                 align=PP_ALIGN.CENTER)
    add_text_box(slide, f"{pricing.get('heston_call', 0):.4f}",
                 6.8, 3.45, 5.4, 0.6,
                 font_size=24, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER)
    diff = pricing.get('difference_pct', 0)
    diff_color = GREEN if diff > 0 else RED
    add_text_box(slide,
                 f"Vol stochastique | Diff : {diff:+.2f}%",
                 6.8, 3.95, 5.4, 0.4,
                 font_size=10, color=diff_color,
                 align=PP_ALIGN.CENTER)

    # Interprétation
    add_card(slide, 0.3, 4.6, 12.1, 1.6)
    add_text_box(slide, "POURQUOI HESTON DIFFÈRE DE BLACK-SCHOLES ?",
                 0.5, 4.7, 11.7, 0.4,
                 font_size=11, bold=True, color=ACCENT)
    add_text_box(slide,
                 "Black-Scholes suppose une volatilité CONSTANTE — irréaliste car la "
                 "volatilité réelle fluctue dans le temps.\n"
                 "Heston modélise la volatilité comme un processus stochastique (CIR), "
                 "capturant le smile de volatilité et les régimes de marché.\n"
                 "Résultat : Heston price mieux les options ITM/OTM et les produits exotiques.",
                 0.5, 5.1, 11.7, 1.0,
                 font_size=11, color=WHITE)

    return slide

def create_portfolio_slide(prs, portfolio_results):
    """Slide optimisation de portefeuille"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_background(slide, prs)

    add_text_box(slide,
                 "OPTIMISATION DE PORTEFEUILLE — MARKOWITZ & BLACK-LITTERMAN",
                 0.3, 0.15, 12, 0.6,
                 font_size=16, bold=True, color=ACCENT)

    line = slide.shapes.add_shape(
        1, Inches(0.3), Inches(0.8), Inches(12.1), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    asset_names = portfolio_results.get('asset_names', [])
    portfolios = portfolio_results.get('portfolios', {})

    add_text_box(slide,
                 f"Actifs : {', '.join(asset_names)}",
                 0.3, 1.0, 12.1, 0.4,
                 font_size=11, color=GRAY,
                 align=PP_ALIGN.CENTER)

    port_labels = {
        'equal_weight': ('Équipondéré', ACCENT),
        'min_variance': ('Variance Min.', GREEN),
        'max_sharpe': ('Sharpe Max.', GOLD),
        'risk_parity': ('Risk Parity', RGBColor(155, 89, 182)),
    }

    card_w = 2.9
    y_start = 1.5

    for i, (key, (label, color)) in enumerate(port_labels.items()):
        if key not in portfolios:
            continue
        p = portfolios[key]
        x = 0.3 + i * (card_w + 0.15)

        add_card(slide, x, y_start, card_w, 3.2)
        add_text_box(slide, label, x, y_start + 0.1, card_w, 0.45,
                     font_size=11, bold=True, color=color,
                     align=PP_ALIGN.CENTER)

        metrics = [
            ('Rendement', f"{p['rendement']:.1f}%",
             GREEN if p['rendement'] > 0 else RED),
            ('Volatilité', f"{p['volatilite']:.1f}%", GOLD),
            ('Sharpe', f"{p['sharpe']:.3f}",
             GREEN if p['sharpe'] > 1 else GOLD),
        ]

        for j, (m_label, m_val, m_color) in enumerate(metrics):
            y_m = y_start + 0.65 + j * 0.6
            add_text_box(slide, m_label, x + 0.1, y_m, card_w - 0.2, 0.3,
                         font_size=9, color=GRAY)
            add_text_box(slide, m_val, x + 0.1, y_m + 0.28,
                         card_w - 0.2, 0.35,
                         font_size=14, bold=True, color=m_color)

        # Allocations
        add_text_box(slide, "Allocations :",
                     x + 0.1, y_start + 2.35, card_w - 0.2, 0.3,
                     font_size=9, bold=True, color=ACCENT)

        for k, (asset, weight) in enumerate(
                zip(asset_names, p['weights'])):
            add_text_box(slide,
                         f"{asset[:12]}: {weight*100:.1f}%",
                         x + 0.1, y_start + 2.65 + k * 0.22,
                         card_w - 0.2, 0.25,
                         font_size=9, color=WHITE)

    # Black-Litterman
    bl = portfolio_results.get('black_litterman', {})
    if bl and bl.get('portfolio'):
        bl_port = bl['portfolio']
        add_card(slide, 0.3, 4.9, 12.1, 1.5)
        add_text_box(slide,
                     "BLACK-LITTERMAN",
                     0.5, 5.0, 4.0, 0.45,
                     font_size=13, bold=True, color=ACCENT)
        add_text_box(slide,
                     f"Rendement : {bl_port['rendement']:.1f}%  |  "
                     f"Volatilité : {bl_port['volatilite']:.1f}%  |  "
                     f"Sharpe : {bl_port['sharpe']:.3f}",
                     0.5, 5.45, 11.7, 0.45,
                     font_size=12, color=WHITE)
        alloc_text = "  |  ".join(
            [f"{a}: {w*100:.1f}%"
             for a, w in zip(asset_names, bl_port['weights'])]
        )
        add_text_box(slide, f"Allocations : {alloc_text}",
                     0.5, 5.9, 11.7, 0.4,
                     font_size=11, color=GOLD)

    return slide

def generate_report(analysis, results, chart_paths, output_filename=None):
    """Génère le rapport PowerPoint complet"""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    asset_name = analysis.get('actifs_detectes', ['Actif Financier'])[0]

    print("Création du rapport PowerPoint...")

    # Slide 1 : Titre
    create_title_slide(prs, analysis, asset_name)
    print("  ✓ Slide titre")

    # Slide 2 : Résumé
    create_summary_slide(prs, analysis)
    print("  ✓ Slide résumé")

    # Slides Black-Scholes
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'black_scholes' in data and data['black_scholes']:
            create_bs_greeks_slide(prs, data['black_scholes'], sheet_name)
            print(f"  ✓ Slide Black-Scholes — {sheet_name}")

    # Slides MC Options
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'mc_options' in data and data['mc_options']:
            create_mc_options_slide(prs, data['mc_options'], sheet_name)
            print(f"  ✓ Slide MC Options — {sheet_name}")

    # Slides Heston
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'heston' in data and data['heston']:
            create_heston_slide(prs, data['heston'], sheet_name)
            print(f"  ✓ Slide Heston — {sheet_name}")

    # Slide Portefeuille
    if 'portfolio_optimization' in results and results['portfolio_optimization']:
        create_portfolio_slide(prs, results['portfolio_optimization'])
        print("  ✓ Slide optimisation portefeuille")

    # Slides signaux de trading
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'signals' in data and data['signals']:
            create_signals_slide(prs, data['signals'], sheet_name)
            print(f"  ✓ Slide signaux — {sheet_name}")

    # Slides métriques de base
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'metrics' in data:
            create_metrics_slide(prs, data['metrics'], sheet_name)
            print(f"  ✓ Slide métriques — {sheet_name}")

    # Slide commentaire IA
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'metrics' in data:
            create_commentary_slide(prs, data['metrics'], analysis, sheet_name)
            print("  ✓ Slide commentaire IA")
            break

    # Slides stress testing textuelles
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'stress' in data:
            create_stress_summary_slide(prs, data['stress'], sheet_name)
            print(f"  ✓ Slide stress test — {sheet_name}")

    # Slides factor model textuelles
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'factors' in data:
            create_factor_summary_slide(prs, data['factors'], sheet_name)
            print(f"  ✓ Slide factor model — {sheet_name}")

    # Slides Monte Carlo textuelles
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'monte_carlo' in data:
            create_monte_carlo_slide(prs, data['monte_carlo'], sheet_name)
            print(f"  ✓ Slide Monte Carlo — {sheet_name}")

    # Slides backtesting textuelles
    for sheet_name, data in results.items():
        if sheet_name in ['correlation_matrix', 'portfolio_optimization']:
            continue
        if 'backtest' in data:
            create_backtest_slide(prs, data['backtest'], sheet_name)
            print(f"  ✓ Slide backtesting — {sheet_name}")

    # Toutes les slides graphiques
    chart_titles = {
        'price_history': 'Évolution Historique des Prix',
        'returns_dist': 'Distribution des Rendements',
        'drawdown': 'Analyse du Drawdown',
        'moving_averages': 'Moyennes Mobiles',
        'rsi': 'RSI — Relative Strength Index',
        'bollinger': 'Bandes de Bollinger',
        'macd': 'MACD',
        'metrics_dashboard': 'Tableau de Bord des Métriques',
        'rolling_sharpe': 'Sharpe Ratio Glissant',
        'rolling_volatility': 'Volatilité Glissante',
        'monte_carlo': 'Simulation Monte Carlo',
        'monte_carlo_stats': 'Statistiques Monte Carlo',
        'backtest_comparison': 'Comparaison des Stratégies',
        'stress_test': 'Stress Testing',
        'factor_model': 'Modèles Factoriels',
        'portfolio_optimization': 'Frontière Efficiente de Markowitz',
        'portfolio_comparison': 'Comparaison des Portefeuilles',
        'correlation_matrix': 'Matrice de Corrélation',
    }

    for chart_path in chart_paths:
        filename = os.path.basename(chart_path).replace('.png', '')
        # Extraire le type depuis le nom du fichier
        parts = filename.split('_', 2)
        chart_type = '_'.join(parts[2:]) if len(parts) >= 3 else filename
        # Chercher le titre correspondant
        title = None
        for key, val in chart_titles.items():
            if key in chart_type:
                title = val
                break
        if title is None:
            title = chart_type.replace('_', ' ').title()

        # Ajouter le nom de l'actif si possible
        if len(parts) >= 2 and parts[1] not in ['', '00']:
            asset_part = parts[1].capitalize()
            if asset_part.lower() not in title.lower():
                title = f"{title} — {asset_part}"

        create_chart_slide(prs, chart_path, title)
        print(f"  ✓ Slide graphique — {title}")

    # Slide conclusion
    create_conclusion_slide(prs, analysis)
    print("  ✓ Slide conclusion")

    # Sauvegarde
    if output_filename is None:
        date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filename = f"rapport_quant_{date_str}.pptx"

    output_path = os.path.join(OUTPUT_DIR, output_filename)
    prs.save(output_path)
    print(f"\n✅ Rapport sauvegardé : {output_path}")
    return output_path